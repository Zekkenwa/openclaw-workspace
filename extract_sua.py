import fitz  # pymupdf
from pptx import Presentation
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')

base = r"C:\Users\chals\Documents\PT SUA"

lines = []
lines.append("="*70)
lines.append("PDF: instalasi udara tekan.pdf")
lines.append("="*70)
doc = fitz.open(base + r"\instalasi udara tekan.pdf")
lines.append(f"Pages: {doc.page_count}\n")
for i, page in enumerate(doc):
    txt = page.get_text().strip()
    if txt:
        lines.append(f"----- PDF PAGE {i+1} -----")
        lines.append(txt)
        lines.append("")
doc.close()

lines.append("="*70)
lines.append("PPTX: instalasi udara tekan.pptx")
lines.append("="*70)
prs = Presentation(base + r"\instalasi udara tekan.pptx")
lines.append(f"Slides: {len(prs.slides)}\n")
for i, slide in enumerate(prs.slides):
    lines.append(f"----- SLIDE {i+1} -----")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                line = "".join(run.text for run in para.runs)
                if line.strip():
                    lines.append(line)
        if shape.has_table:
            for row in shape.table.rows:
                cells = [c.text for c in row.cells]
                lines.append(" | ".join(cells))
    lines.append("")

output = "\n".join(lines)
with open(r"C:\Users\chals\.openclaw\workspace\sua_extracted.txt", "w", encoding="utf-8") as f:
    f.write(output)
print("Done! Written to sua_extracted.txt")